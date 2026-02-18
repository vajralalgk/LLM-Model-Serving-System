"""
PowerPoint Presentation Generator for LLM Model Serving System
Author: Gopi Krishna Vajrala

Generates a professional presentation covering the system architecture,
design decisions, and implementation details.

Usage: python scripts/generate_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Color scheme
DARK_BLUE = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xDE)
LIGHT_BLUE = RGBColor(0xE8, 0xF4, 0xFD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT_GRAY = RGBColor(0x95, 0x95, 0x95)
GREEN = RGBColor(0x27, 0xAE, 0x60)
ORANGE = RGBColor(0xF3, 0x9C, 0x12)
RED = RGBColor(0xE7, 0x4C, 0x3C)


def add_background(slide, color=DARK_BLUE):
    """Add a solid background color to a slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title_text, subtitle_text=None):
    """Add a styled title bar to the top of a slide."""
    # Title shape
    left = Inches(0.5)
    top = Inches(0.3)
    width = Inches(9)
    height = Inches(0.8)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(14)
        p2.font.color.rgb = LIGHT_GRAY


def add_bullet_content(slide, items, left=0.5, top=1.5, width=8.5, font_size=16):
    """Add bulleted content to a slide."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i > 0:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]

        if isinstance(item, tuple):
            p.text = item[0]
            p.font.size = Pt(font_size)
            p.font.bold = item[1] if len(item) > 1 else False
            p.font.color.rgb = item[2] if len(item) > 2 else WHITE
        else:
            p.text = f"  {item}"
            p.font.size = Pt(font_size)
            p.font.color.rgb = WHITE
        p.space_after = Pt(8)


def add_two_column(slide, left_items, right_items, left_title="", right_title=""):
    """Add two-column layout."""
    # Left column title
    if left_title:
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(4), Inches(0.5))
        p = txBox.text_frame.paragraphs[0]
        p.text = left_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

    # Right column title
    if right_title:
        txBox = slide.shapes.add_textbox(Inches(5.2), Inches(1.3), Inches(4), Inches(0.5))
        p = txBox.text_frame.paragraphs[0]
        p.text = right_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

    add_bullet_content(slide, left_items, left=0.5, top=1.9, width=4, font_size=13)
    add_bullet_content(slide, right_items, left=5.2, top=1.9, width=4, font_size=13)


def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ========== SLIDE 1: Title Slide ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    add_background(slide, DARK_BLUE)

    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Real-Time LLM Model Serving System"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = "Production-Grade Personalization Ranking Engine"
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT_BLUE
    p2.alignment = PP_ALIGN.CENTER

    # Accent line
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(3.5), Inches(4), Inches(0.05))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()

    # Author
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(1.5))
    tf2 = txBox2.text_frame
    p3 = tf2.paragraphs[0]
    p3.text = "Author: Gopi Krishna Vajrala"
    p3.font.size = Pt(18)
    p3.font.color.rgb = WHITE
    p3.alignment = PP_ALIGN.CENTER
    p4 = tf2.add_paragraph()
    p4.text = "Production Engineering | Distributed Systems | ML Infrastructure"
    p4.font.size = Pt(12)
    p4.font.color.rgb = LIGHT_GRAY
    p4.alignment = PP_ALIGN.CENTER

    # ========== SLIDE 2: Problem Statement ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Problem Statement", "Why build this system?")
    add_bullet_content(slide, [
        ("The Challenge:", True, ACCENT_BLUE),
        "Modern content platforms need real-time, personalized ranking",
        "Traditional recommendation systems have high latency (500ms+)",
        "LLM-based ranking provides semantic understanding but is compute-heavy",
        "",
        ("Our Solution:", True, GREEN),
        "Production-grade LLM serving with <120ms p95 latency",
        "Dynamic batching for optimal GPU utilization",
        "5,000+ RPS with horizontal auto-scaling",
        "99.99% availability with circuit breakers & graceful degradation",
    ])

    # ========== SLIDE 3: System Architecture ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "System Architecture", "End-to-end request flow")
    add_bullet_content(slide, [
        ("Request Flow:", True, ACCENT_BLUE),
        "1. Client sends POST /rank with user context + candidate titles",
        "2. Ingress: TLS termination, rate limiting, load balancing",
        "3. Auth Layer: JWT/API key validation",
        "4. Dynamic Batcher: Collects requests (10-20ms window)",
        "5. Model Server: Batch inference on GPU/CPU",
        "6. Ranking Engine: Cosine similarity + personalization",
        "7. De-batch: Individual responses via async futures",
        "8. Response: Ranked titles with scores in <120ms",
    ], font_size=14)

    # ========== SLIDE 4: Core Components ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Core Components")
    add_two_column(slide,
        left_items=[
            ("FastAPI Application", True, ACCENT_BLUE),
            "Async-first, high-performance Python",
            "Pydantic validation for type safety",
            "Auto-generated OpenAPI docs",
            "",
            ("Dynamic Batcher", True, ACCENT_BLUE),
            "10-20ms collection window",
            "Max 32 requests per batch",
            "3-5x throughput improvement",
            "Async Future-based delivery",
        ],
        right_items=[
            ("Model Server", True, ACCENT_BLUE),
            "Sentence-transformer embeddings",
            "Cosine similarity scoring",
            "User personalization signals",
            "",
            ("Model Registry", True, ACCENT_BLUE),
            "Multi-version management",
            "Hot-swap without restart",
            "Blue/Green deployment ready",
            "Lightweight fallback model",
        ],
        left_title="Application Layer",
        right_title="ML Layer",
    )

    # ========== SLIDE 5: Performance Engineering ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Performance Engineering", "Meeting strict SLA requirements")
    add_bullet_content(slide, [
        ("Performance Targets:", True, ACCENT_BLUE),
        "  p95 Latency: < 120ms (target) | ~23ms achieved with lightweight model",
        "  Throughput: 5,000 RPS sustained with horizontal scaling",
        "  Availability: 99.99% with 3+ replicas and circuit breakers",
        "",
        ("Key Optimizations:", True, GREEN),
        "  Dynamic Batching: Combines requests into single GPU forward pass",
        "  Connection Pooling: Reuse connections to reduce overhead",
        "  ORJson Serialization: 10x faster than standard json",
        "  Async I/O: Non-blocking operations throughout the stack",
        "  Thread Pool: CPU-bound inference runs in executor threads",
        "",
        ("Auto-Scaling Strategy:", True, ORANGE),
        "  HPA: CPU 70%, Memory 80%, Custom RPS metric",
        "  Scale Up: Aggressive (4 pods/min or 100%/min)",
        "  Scale Down: Conservative (1 pod/min, 5min stabilization)",
    ], font_size=13)

    # ========== SLIDE 6: Resilience & Reliability ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Resilience & Reliability", "Designed for 99.99% availability")
    add_two_column(slide,
        left_items=[
            ("Circuit Breaker", True, ACCENT_BLUE),
            "3-state: CLOSED -> OPEN -> HALF_OPEN",
            "5 failures threshold to OPEN",
            "30s recovery timeout",
            "Prevents cascade failures",
            "",
            ("Retry Mechanism", True, ACCENT_BLUE),
            "Exponential backoff with jitter",
            "Configurable max retries (3)",
            "Prevents thundering herd",
        ],
        right_items=[
            ("Timeout Enforcement", True, ACCENT_BLUE),
            "5s default request timeout",
            "Prevents hanging requests",
            "Configurable per-endpoint",
            "",
            ("Graceful Shutdown", True, ACCENT_BLUE),
            "SIGTERM handler",
            "Drain batcher queue",
            "Complete in-flight requests",
            "30s termination grace period",
            "",
            ("Chaos Testing", True, RED),
            "Pod kill under load",
            "Resource pressure simulation",
            "Service endpoint stress test",
        ],
        left_title="Fault Tolerance",
        right_title="Stability Patterns",
    )

    # ========== SLIDE 7: Security ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Security & Governance", "Defense-in-depth approach")
    add_bullet_content(slide, [
        ("Authentication:", True, ACCENT_BLUE),
        "  JWT Bearer tokens with configurable expiration",
        "  API Key authentication for service-to-service calls",
        "  Public paths exempted (health, metrics)",
        "",
        ("Request Protection:", True, GREEN),
        "  Rate limiting: 100 req/min per IP (configurable)",
        "  Request size limit: 1MB maximum payload",
        "  Input validation: Pydantic schema enforcement",
        "  Max 500 candidate titles per request",
        "",
        ("Infrastructure Security:", True, ORANGE),
        "  Non-root Docker containers",
        "  TLS/SSL via Ingress",
        "  Kubernetes RBAC",
        "  Secrets management via K8s Secrets",
        "  Audit logging for all requests",
    ], font_size=14)

    # ========== SLIDE 8: Observability ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Observability Stack", "Full visibility into system behavior")
    add_two_column(slide,
        left_items=[
            ("Prometheus Metrics", True, ACCENT_BLUE),
            "p50/p95/p99 latency histograms",
            "Request rate by status",
            "Error rate by type",
            "Batch size distribution",
            "Queue depth gauge",
            "Circuit breaker state",
            "Model version traffic",
        ],
        right_items=[
            ("Structured Logging", True, ACCENT_BLUE),
            "JSON format via structlog",
            "Request correlation IDs",
            "Timing in every log entry",
            "",
            ("Distributed Tracing", True, ACCENT_BLUE),
            "OpenTelemetry integration",
            "End-to-end request tracing",
            "Jaeger/Zipkin compatible",
            "",
            ("Grafana Dashboards", True, GREEN),
            "Pre-built dashboard configs",
            "Real-time RPS monitoring",
            "Latency SLA tracking",
        ],
        left_title="Metrics",
        right_title="Logging & Tracing",
    )

    # ========== SLIDE 9: Deployment Strategy ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Deployment Strategy", "Zero-downtime deployments on Kubernetes")
    add_bullet_content(slide, [
        ("Blue/Green Deployment:", True, ACCENT_BLUE),
        "  Two identical deployments (blue + green) running simultaneously",
        "  Traffic switched via Kubernetes Service selector patch",
        "  Instant rollback by switching selector back",
        "  Zero downtime during model updates",
        "",
        ("Kubernetes Infrastructure:", True, GREEN),
        "  3 replicas minimum (HA guarantee)",
        "  HPA: Auto-scale to 20 replicas under load",
        "  Rolling updates with maxSurge=1, maxUnavailable=0",
        "  Startup/Liveness/Readiness probes configured",
        "  Pre-stop hook: 10s sleep for connection draining",
        "",
        ("GPU Support:", True, ORANGE),
        "  Dedicated Dockerfile.gpu with NVIDIA CUDA 12.1",
        "  GPU node pool support (GKE/EKS/AKS)",
        "  CPU fallback with lightweight ranking model",
    ], font_size=13)

    # ========== SLIDE 10: Testing Strategy ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Testing & Validation", "Comprehensive quality assurance")
    add_two_column(slide,
        left_items=[
            ("Unit Tests", True, ACCENT_BLUE),
            "API endpoint testing",
            "Model server validation",
            "Circuit breaker states",
            "Dynamic batcher behavior",
            "Schema validation",
            "",
            ("Integration Tests", True, ACCENT_BLUE),
            "Full request pipeline",
            "Auth flow validation",
            "Error handling paths",
        ],
        right_items=[
            ("Load Tests", True, GREEN),
            "Locust: Web UI + scripted",
            "k6: CLI-based with thresholds",
            "Target: 5000 RPS / <120ms p95",
            "",
            ("Chaos Tests", True, RED),
            "Pod kill under load",
            "Resource pressure simulation",
            "Network partition testing",
            "Recovery time measurement",
            "",
            ("CI/CD Ready", True, ORANGE),
            "pytest with coverage reports",
            "JUnit XML output for CI",
            "Makefile automation",
        ],
        left_title="Functional Testing",
        right_title="Non-Functional Testing",
    )

    # ========== SLIDE 11: Tech Stack ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Technology Stack")
    add_two_column(slide,
        left_items=[
            ("Backend Framework", True, ACCENT_BLUE),
            "  Python 3.11 + FastAPI",
            "  Uvicorn + Gunicorn (ASGI)",
            "  Pydantic v2 (validation)",
            "",
            ("ML / AI", True, ACCENT_BLUE),
            "  PyTorch (inference engine)",
            "  Sentence-Transformers (embeddings)",
            "  NumPy (vector operations)",
            "",
            ("Infrastructure", True, ACCENT_BLUE),
            "  Docker (multi-stage builds)",
            "  Kubernetes (orchestration)",
            "  NGINX Ingress (load balancing)",
            "  Redis (rate limiting cache)",
        ],
        right_items=[
            ("Observability", True, GREEN),
            "  Prometheus (metrics)",
            "  Grafana (dashboards)",
            "  OpenTelemetry (tracing)",
            "  structlog (JSON logging)",
            "",
            ("Security", True, GREEN),
            "  python-jose (JWT)",
            "  SlowAPI (rate limiting)",
            "",
            ("Testing", True, GREEN),
            "  pytest + pytest-asyncio",
            "  Locust (load testing)",
            "  k6 (performance testing)",
            "  Chaos scripts (reliability)",
        ],
        left_title="Core Stack",
        right_title="Supporting Stack",
    )

    # ========== SLIDE 12: Future Roadmap ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)
    add_title_bar(slide, "Future Roadmap", "Elite-level extensions planned")
    add_bullet_content(slide, [
        ("Phase 2 - Advanced Serving:", True, ACCENT_BLUE),
        "  Multi-region deployment with geo-routing",
        "  Canary deployment with A/B model testing",
        "  Token streaming support for real-time responses",
        "  ONNX Runtime optimization for faster inference",
        "",
        ("Phase 3 - Intelligence:", True, GREEN),
        "  Cost-aware routing (small vs. large model selection)",
        "  Shadow traffic testing for new models",
        "  Online learning from user feedback",
        "  Feature store integration for rich user profiles",
        "",
        ("Phase 4 - Scale:", True, ORANGE),
        "  Multi-cloud deployment (GKE + EKS + AKS)",
        "  Edge deployment for ultra-low latency",
        "  Custom CUDA kernels for model optimization",
        "  50,000+ RPS target with GPU cluster",
    ], font_size=14)

    # ========== SLIDE 13: Thank You ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, DARK_BLUE)

    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(3.5), Inches(3), Inches(0.05))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()

    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(2))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "Gopi Krishna Vajrala"
    p2.font.size = Pt(22)
    p2.font.color.rgb = ACCENT_BLUE
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf2.add_paragraph()
    p3.text = "Real-Time LLM Model Serving System"
    p3.font.size = Pt(14)
    p3.font.color.rgb = LIGHT_GRAY
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf2.add_paragraph()
    p4.text = "Production-Grade | Kubernetes-Native | GPU-Optimized"
    p4.font.size = Pt(12)
    p4.font.color.rgb = LIGHT_GRAY
    p4.alignment = PP_ALIGN.CENTER

    # Save
    output_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "docs", "LLM_Ranking_Service_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")
    return output_path


if __name__ == "__main__":
    create_presentation()
